import os
from dataclasses import dataclass
from typing import Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from pandas import DataFrame


@dataclass
class Case:
    providerInvoiceNumber: str
    filePrefix: str
    workflowNumber: Optional[str]


def remove_empty_invoice_id_rows_from_df(df: DataFrame) -> DataFrame:
    return df[df.Invoice_ID.notna() & (df.Invoice_ID != "")]


def find_data_sheet(excel_path: str) -> str:
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"File does not exist: {excel_path}")

    try:
        book = load_workbook(excel_path, read_only=True)
    except Exception as e:
        raise Exception(f"Nie można otworzyć pliku Excel: {e}")

    required_columns = {'Lp', 'Invoice_ID', 'WF'}

    try:
        for sheet_name in book.sheetnames:
            try:
                df = pd.read_excel(excel_path, sheet_name=sheet_name, engine="openpyxl", nrows=0)
                if required_columns.issubset(set(df.columns)):
                    return sheet_name
            except Exception:
                continue

        raise Exception(
            f"Nie znaleziono zakładki z wymaganymi kolumnami (Lp, Invoice_ID, WF).\n"
        )
    finally:
        book.close()

def read_input_cases(path: str) -> tuple[list[Case], DataFrame, str]:
    cases = []
    try:
        sheet_name = find_data_sheet(path)
        df = pd.read_excel(path, sheet_name=sheet_name, engine='openpyxl')
        df_without_none = remove_empty_invoice_id_rows_from_df(df)
    except FileNotFoundError:
        raise Exception('Failed to open input cases file, file does not exist.')

    for i, (row_index, row_cells) in enumerate(df_without_none.iterrows()):
        if pd.isna(row_cells['Lp']):
            raise Exception(f'Row #{i + 1} does not contain an LP number.')
        cases.append(convert_row_to_case(row_cells))
    return cases, df_without_none, sheet_name


def is_excel_file_open(excel_path: str) -> bool:
    try:
        with open(excel_path, 'r+b'):
            return False
    except PermissionError:
        return True
    except FileNotFoundError:
        return False


def convert_row_to_case(row: pd.Series) -> Case:
    return Case(
        providerInvoiceNumber=str(row['Invoice_ID']),
        filePrefix=str(row['Lp']),
        workflowNumber=get_cell_optional(row, 'WF'))


def get_cell_optional(row: pd.Series, column: str) -> Optional[str]:
    if pd.isna(row[column]):
        return None
    return str(row[column])


def update_excel_df(main_df: DataFrame, found_invoices: list[str], excel_path: str, sheet_name: str):
    main_df["Comments"] = main_df["Invoice_ID"].apply(
        lambda x: "File was found" if x in found_invoices else "File wasn't found")
    try:
        book = load_workbook(excel_path)
        ws = book[sheet_name]
        ws.delete_rows(1, ws.max_row)

        for r in dataframe_to_rows(main_df, index=False, header=True):
            ws.append(r)
        book.save(excel_path)
        book.close()

    except Exception as e:
        raise Exception(f"Failed to save excel file {e}")
